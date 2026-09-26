"""
Exporta as cifras extraídas para Word, PDF, texto e Excel.
Cada função recebe a lista de músicas (dicts) e devolve o arquivo em bytes.
Opções: colunas (1 ou 2) e tamanho da fonte (None = automático).
Uma música pode ter "Colunas" próprio (1 ou 2), que tem prioridade sobre a opção geral.
"""

import io
from itertools import zip_longest

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from layout import ENTRELINHA, PT_MM, montar, paginar, preparar

# Página A4 (mm)
PAGINA_L, PAGINA_A = 210, 297
MARGEM_LADO, MARGEM_TOPO = 15, 12
ESPACO_COLUNAS = 8
CABECALHO = 20  # altura do título da música
COR_ACORDE = (150, 30, 30)
LARGURA_TXT = 48  # caracteres por coluna no .txt com 2 colunas


def _nome(item):
    """'Música (Artista)' ou um aviso quando a extração falhou"""
    if item["Status"] != "OK":
        return "(cifra não extraída)"
    return f"{item['Música']} ({item['Artista']})" if item["Artista"] else item["Música"]


def _cifra(item):
    return item["Cifra"] if item["Status"] == "OK" else f"Não foi possível extrair: {item['Erro']}\n{item['URL']}"


def _colunas(item, padrao):
    return item.get("Colunas") or padrao


def _area_colunas(colunas, folga=0):
    """Largura de cada coluna e altura disponível para a cifra (mm)"""
    largura = (PAGINA_L - 2 * MARGEM_LADO - ESPACO_COLUNAS * (colunas - 1)) / colunas
    altura = PAGINA_A - 2 * MARGEM_TOPO - CABECALHO - folga
    return largura, altura


# ------------------------------------------------------------------- TXT
def exportar_txt(dados, colunas=1, tamanho=None):
    blocos = ["CIFRAS DA MISSA", ""]
    for item in dados:
        blocos += ["=" * 60, item["Posição"].upper(), _nome(item), "=" * 60, ""]
        if _colunas(item, colunas) == 1:
            blocos.append(_cifra(item))
        else:
            unidades, _ = preparar(_cifra(item), LARGURA_TXT)
            esquerda, *direita = paginar(unidades, 10**6, 2)[0]
            direita = direita[0] if direita else []
            for a, b in zip_longest(esquerda, direita, fillvalue=("vazia", "")):
                blocos.append(f"{a[1]:<{LARGURA_TXT}}  |  {b[1]}".rstrip())
        blocos += ["", ""]
    # utf-8-sig para o Bloco de Notas reconhecer os acentos
    return "\n".join(blocos).encode("utf-8-sig")


# ------------------------------------------------------------------- Word
def _definir_colunas(secao, n):
    cols = secao._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        secao._sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(ESPACO_COLUNAS * 56.7)))  # mm -> twips
    cols.set(qn("w:sep"), "1" if n > 1 else "0")               # linha entre as colunas


def _configurar_pagina(secao):
    secao.page_width, secao.page_height = Mm(PAGINA_L), Mm(PAGINA_A)
    secao.left_margin = secao.right_margin = Mm(MARGEM_LADO)
    secao.top_margin = secao.bottom_margin = Mm(MARGEM_TOPO)


def _paragrafo(doc, texto, tamanho, negrito=False, cor=None, fonte="Courier New"):
    p = doc.add_paragraph()
    formato = p.paragraph_format
    formato.space_before = formato.space_after = Pt(0)
    formato.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    formato.line_spacing = Pt(tamanho * ENTRELINHA)
    run = p.add_run(texto)
    run.font.name = fonte
    run.font.size = Pt(tamanho)
    run.bold = negrito
    if cor:
        run.font.color.rgb = RGBColor(*cor)
    return p


def exportar_docx(dados, colunas=1, tamanho=None):
    doc = Document()
    _configurar_pagina(doc.sections[0])
    for i, item in enumerate(dados):
        n_colunas = _colunas(item, colunas)
        # Folga de 10 mm: o Word calcula alturas um pouco diferente do PDF
        largura, altura = _area_colunas(n_colunas, folga=10)
        secao = doc.sections[-1] if i == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
        _definir_colunas(secao, 1)
        _paragrafo(doc, item["Posição"], 16, negrito=True, fonte="Arial").paragraph_format.space_after = Pt(2)
        _paragrafo(doc, _nome(item), 12, fonte="Arial").paragraph_format.space_after = Pt(10)

        # Conteúdo numa seção contínua com as colunas; quebras de coluna calculadas pelo layout
        _definir_colunas(doc.add_section(WD_SECTION.CONTINUOUS), n_colunas)
        t, paginas = montar(_cifra(item), n_colunas, largura, altura, tamanho)
        todas = [coluna for pagina in paginas for coluna in pagina]
        for n, coluna in enumerate(todas):
            ultimo = None
            for tipo, texto in coluna:
                negrito = tipo in ("acorde", "marcador")
                ultimo = _paragrafo(doc, texto, t, negrito, COR_ACORDE if tipo == "acorde" else None)
            if n < len(todas) - 1:
                if ultimo is None:
                    ultimo = _paragrafo(doc, "", t)
                ultimo.add_run().add_break(WD_BREAK.COLUMN if n_colunas > 1 else WD_BREAK.PAGE)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------- PDF
def _latin1(texto):
    """As fontes padrão do PDF só aceitam latin-1 (que cobre os acentos do português)"""
    trocas = {"—": "-", "–": "-", "“": '"', "”": '"', "‘": "'", "’": "'", "…": "..."}
    for antigo, novo in trocas.items():
        texto = texto.replace(antigo, novo)
    return texto.encode("latin-1", "replace").decode("latin-1")


def exportar_pdf(dados, colunas=1, tamanho=None):
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(False)
    topo = MARGEM_TOPO + CABECALHO

    for item in dados:
        n_colunas = _colunas(item, colunas)
        largura, altura = _area_colunas(n_colunas)
        t, paginas = montar(_cifra(item), n_colunas, largura, altura, tamanho)
        altura_linha = t * ENTRELINHA * PT_MM

        for n, pagina in enumerate(paginas):
            pdf.add_page()
            # Cabeçalho
            pdf.set_text_color(0, 0, 0)
            pdf.set_xy(MARGEM_LADO, MARGEM_TOPO)
            pdf.set_font("Helvetica", "B", 16)
            continuacao = f"  (continuação {n + 1}/{len(paginas)})" if n else ""
            pdf.cell(0, 8, _latin1(item["Posição"] + continuacao), new_x="LMARGIN", new_y="NEXT")
            pdf.set_x(MARGEM_LADO)
            pdf.set_font("Helvetica", "", 12)
            pdf.cell(0, 7, _latin1(_nome(item)))
            pdf.set_draw_color(180, 180, 180)
            pdf.line(MARGEM_LADO, topo - 3, PAGINA_L - MARGEM_LADO, topo - 3)

            # Colunas
            for c, coluna in enumerate(pagina):
                x = MARGEM_LADO + c * (largura + ESPACO_COLUNAS)
                if c:
                    divisa = x - ESPACO_COLUNAS / 2
                    pdf.line(divisa, topo, divisa, topo + len(coluna) * altura_linha)
                for k, (tipo, texto) in enumerate(coluna):
                    if tipo == "vazia":
                        continue
                    pdf.set_font("Courier", "B" if tipo in ("acorde", "marcador") else "", t)
                    pdf.set_text_color(*(COR_ACORDE if tipo == "acorde" else (0, 0, 0)))
                    pdf.text(x, topo + (k + 0.8) * altura_linha, _latin1(texto))

    return bytes(pdf.output())


# ------------------------------------------------------------------- Excel
def exportar_xlsx(dados, colunas=1, tamanho=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cifras"
    ws.append(["Posição", "Música", "Artista", "Cifra", "URL"])
    for item in dados:
        ws.append([item["Posição"], item["Música"], item["Artista"], _cifra(item), item["URL"]])

    # Estilos
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    lado = Side(style="thin")
    border = Border(left=lado, right=lado, top=lado, bottom=lado)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for coluna, largura in zip("ABCDE", (20, 30, 25, 80, 40)):
        ws.column_dimensions[coluna].width = largura

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        row[3].font = Font(name="Courier New", size=10)
        # Excel não ajusta a altura sozinho; 13 pontos por linha da cifra (máx. 409)
        ws.row_dimensions[row[0].row].height = min(409, 13 * (str(row[3].value).count("\n") + 1))

    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


EXPORTADORES = {
    "docx": exportar_docx,
    "pdf": exportar_pdf,
    "txt": exportar_txt,
    "xlsx": exportar_xlsx,
}
